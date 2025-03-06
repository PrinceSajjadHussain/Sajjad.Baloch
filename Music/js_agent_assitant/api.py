import requests
import re
import json
from docx import Document
import google.generativeai as genai

# === Step 1: Extract CV Content from PDF via Landing.ai API ===
landing_url = "https://api.va.landing.ai/v1/tools/agentic-document-analysis"
landing_headers = {
    "Authorization": "Basic bDQwbWQyMGZiNDcwaml2M21kdjVkOlJvdHpzbG1WeW9idmJtMlpsVEs4ODJjSWs1eEE0U2ZY",
}

with open("cv.pdf", "rb") as file:
    files = {"pdf": file}  # Use the "pdf" key for PDF files.
    landing_response = requests.post(landing_url, files=files, headers=landing_headers)

if landing_response.status_code == 200:
    data = landing_response.json().get("data", {})
    # The API returns markdown text – you can override it if needed.
    extracted_text = data.get("markdown", "")
else:
    print("Error extracting document:", landing_response.status_code, landing_response.text)
    exit()

# (Optional) Uncomment to override with your updated CV text:
"""
extracted_text = '''Description
The image is a logo consisting of a simple design. It features a circle with two letters inside: "S" and "B". These letters are separated by a vertical line. The entire design is in a dark blue color. The circle is thin and evenly drawn, enclosing the letters and line neatly. There are no additional elements, text, or colors present in the logo. <!-- figure ... -->
Sajjad Baloch
Python Developer <!-- title ... -->
Contact Information
- **Location**: Karachi, 75500 Sindh
- **Phone**: 03048467517
- **Email**: sajjad.aiengineer@gmail.com <!-- key_value ... -->
PROFESSIONAL SUMMARY
- [ ] Detail-oriented and highly collaborative Python/AI Engineer ... <!-- text ... -->
Skills Overview
The document is a form listing various skills ... 
Skills
- **Python**
  - Proficiency: 5 bars
...
Technical Profile
- [ ] Postman
- [ ] SQL
- [ ] Linux
- [ ] Git
- [ ] Docker <!-- text ... -->'''
"""

# === Step 2: Use Gemini LLM to Determine Best Formatting ===
# Replace with your actual Gemini API key.
GEMINI_API_KEY = "YOUR_GEMINI_API_KEY"
genai.configure(api_key=GEMINI_API_KEY)
model = genai.GenerativeModel("gemini-2.0-flash")

# Construct a prompt instructing the LLM to format the CV.
prompt = (
    "You are an expert CV formatter. Analyze the following extracted CV text and produce a "
    "JSON object representing the best formatted CV document. The JSON object should have section "
    "headings as keys and the corresponding content as values. Use the following JSON format exactly:\n\n"
    "{\n"
    '  "Section Heading 1": "Content...",\n'
    '  "Section Heading 2": "Content...",\n'
    "  ...\n"
    "}\n\n"
    "Extracted CV Text:\n"
    f"{extracted_text}\n\n"
    "Return only valid JSON."
)

gemini_response = model.generate_content(
    prompt=prompt,
    temperature=0.7,
    max_output_tokens=1500
)
gemini_output = gemini_response.result.strip()

# Attempt to parse the Gemini response as JSON.
try:
    formatted_cv = json.loads(gemini_output)
except Exception as e:
    print("Error parsing Gemini output as JSON:", e)
    print("Gemini output:", gemini_output)
    exit()

# === Step 3: Build the DOCX Document from Gemini’s Formatted CV ===
document = Document()

# Simple heuristic: if a heading is in all uppercase, use level 1; else level 2.
def get_heading_level(heading):
    return 1 if heading.isupper() else 2

for heading, content in formatted_cv.items():
    level = get_heading_level(heading)
    document.add_heading(heading, level=level)
    if content:
        document.add_paragraph(content)

output_filename = "formatted_cv.docx"
document.save(output_filename)
print(f"Formatted DOCX file has been created and saved as {output_filename}")

# import requests
# import re
# import json
# from docx import Document

# url = "https://api.va.landing.ai/v1/tools/agentic-document-analysis"
# headers = {
#     "Authorization": "Basic bDQwbWQyMGZiNDcwaml2M21kdjVkOlJvdHpzbG1WeW9idmJtMlpsVEs4ODJjSWs1eEE0U2ZY",
# }

# # Send the PDF file to the API.
# with open("cv.pdf", "rb") as file:
#     files = {"pdf": file}  # use the "pdf" key for PDF files.
#     response = requests.post(url, files=files, headers=headers)

# if response.status_code == 200:
#     # Extract the markdown text from the response.
#     data = response.json().get("data", {})
#     markdown_text = data.get("markdown", "")
    
#     # Split markdown text into sections based on lines starting with "#"
#     sections = re.split(r'\n(?=#)', markdown_text)
#     structured_data = {}
    
#     for section in sections:
#         lines = section.splitlines()
#         if not lines:
#             continue
#         # The first line is assumed to be the heading.
#         heading = lines[0].strip()
#         content = "\n".join(lines[1:]).strip()
#         structured_data[heading] = content

#     # Create a DOCX document.
#     document = Document()

#     def parse_heading(key):
#         """
#         Parse the heading string to determine the level based on the number of '#' characters.
#         Returns (level, text) where level is an integer and text is the heading without '#' characters.
#         """
#         match = re.match(r'^(#+)\s*(.*)', key)
#         if match:
#             level = len(match.group(1))
#             text = match.group(2).strip()
#             return level, text
#         return 0, key

#     # Add sections to the DOCX document.
#     for key, value in structured_data.items():
#         level, heading_text = parse_heading(key)
#         if level > 0:
#             document.add_heading(heading_text, level=level)
#         else:
#             document.add_paragraph(key)
#         if value:
#             document.add_paragraph(value)

#     # Save the DOCX file.
#     output_filename = "output.docx"
#     document.save(output_filename)
#     print(f"DOCX file has been created and saved as {output_filename}")
# else:
#     print("Error:", response.status_code, response.text)


# # import requests
# # import re
# # import json

# # url = "https://api.va.landing.ai/v1/tools/agentic-document-analysis"
# # headers = {
# #     "Authorization": "Basic bDQwbWQyMGZiNDcwaml2M21kdjVkOlJvdHpzbG1WeW9idmJtMlpsVEs4ODJjSWs1eEE0U2ZY",
# # }

# # # Open the PDF file and send it to the API.
# # with open("cv.pdf", "rb") as file:
# #     files = {"pdf": file}  # use "pdf" key for PDF files.
# #     response = requests.post(url, files=files, headers=headers)

# # # Ensure the request was successful.
# # if response.status_code == 200:
# #     # Get the API response data.
# #     data = response.json().get("data", {})
# #     markdown_text = data.get("markdown", "")
    
# #     # Use a regular expression to split sections based on markdown headers.
# #     # Here we assume sections start with either a primary (#) or secondary (##) heading.
# #     # Adjust the regex pattern based on your desired granularity.
# #     sections = re.split(r'\n(?=#)', markdown_text)
    
# #     structured_data = {}
# #     current_section = None

# #     for section in sections:
# #         # Find the heading line.
# #         lines = section.splitlines()
# #         if not lines:
# #             continue

# #         # Assume the first line is the heading.
# #         heading = lines[0].strip()
# #         content = "\n".join(lines[1:]).strip()

# #         # If the heading already exists, you can decide to merge or keep separate.
# #         structured_data[heading] = content

# #     # Print the structured output as formatted JSON.
# #     print(json.dumps(structured_data, indent=2))
# # else:
# #     print("Error:", response.status_code, response.text)


# # # import requests

# # # url = "https://api.va.landing.ai/v1/tools/agentic-document-analysis"
# # # files = {
# # #     "pdf": open("cv.pdf", "rb")
# # # }
# # # headers = {
# # #   "Authorization": "Basic bDQwbWQyMGZiNDcwaml2M21kdjVkOlJvdHpzbG1WeW9idmJtMlpsVEs4ODJjSWs1eEE0U2ZY",
# # # }
# # # response = requests.post(url, files=files, headers=headers)

# # # print(response.json())


